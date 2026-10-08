"""Work out a template's visual style by measuring its slides.

A client hands over a .pptx and nothing else. It may be a real template, an
old deck, or an export from another tool — so nothing here relies on layouts,
placeholders or any standard structure existing. Instead every shape on every
slide is inspected for position, font and color, and the dominant values win.

A .pptx is a ZIP of XML; `python-pptx` parses it into objects, so most of this
is walking `slide.shapes`. The one thing it doesn't expose is the theme color
scheme, so `ppt/theme/theme1.xml` is read directly with lxml — needed because
shapes routinely reference a theme color (`accent1`) instead of a literal hex.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

# python-pptx's MSO_THEME_COLOR names -> the keys used in theme1.xml
_THEME_COLOR_NAMES = {
    "DARK_1": "dk1", "LIGHT_1": "lt1", "DARK_2": "dk2", "LIGHT_2": "lt2",
    "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
    "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
    "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink",
}


@dataclass
class MeasuredStyle:
    """What could be read off a template. Every field may be None — callers
    fall back to their own defaults for anything that wasn't found."""

    slide_width_in: float
    slide_height_in: float
    heading_font: str | None = None
    body_font: str | None = None
    heading_pt: float | None = None
    body_pt: float | None = None
    small_pt: float | None = None
    colors: dict[str, str] = field(default_factory=dict)
    margin_in: float | None = None
    content_right_in: float | None = None
    title_top_in: float | None = None
    content_bottom_in: float | None = None
    logo: bytes | None = None
    logo_box_in: tuple[float, float, float, float] | None = None
    background_image: bytes | None = None
    footer: str | None = None
    notes: list[str] = field(default_factory=list)


def _hex(rgb) -> str:
    return str(rgb).upper().lstrip("#")


def _colorfulness(hex_value: str) -> int:
    """How saturated a color is. Greys, black and white score ~0, which is how
    an accent color is told apart from text and background colors."""
    try:
        r, g, b = (int(hex_value[i : i + 2], 16) for i in (0, 2, 4))
    except (ValueError, IndexError):
        return 0
    return max(r, g, b) - min(r, g, b)


def _theme_palette(prs: Presentation) -> dict[str, str]:
    """Read the theme's color scheme and fonts out of theme1.xml."""
    palette: dict[str, str] = {}
    for part in prs.part.package.iter_parts():
        if "theme" not in str(part.partname):
            continue
        root = etree.fromstring(part.blob)
        scheme = root.find(".//a:clrScheme", _A)
        if scheme is not None:
            for child in scheme:
                name = etree.QName(child).localname
                srgb = child.find("a:srgbClr", _A)
                sys = child.find("a:sysClr", _A)
                if srgb is not None:
                    palette[name] = srgb.get("val", "").upper()
                elif sys is not None and sys.get("lastClr"):
                    palette[name] = sys.get("lastClr").upper()
        fonts = root.find(".//a:fontScheme", _A)
        if fonts is not None:
            for which, key in (("majorFont", "_major"), ("minorFont", "_minor")):
                el = fonts.find(f"a:{which}/a:latin", _A)
                if el is not None and el.get("typeface"):
                    palette[key] = el.get("typeface")
        break
    return palette


def _run_color(run, palette: dict[str, str]) -> str | None:
    """A run's color, resolving a theme reference to a literal hex."""
    try:
        color = run.font.color
        if color is None or color.type is None:
            return None
        if str(color.type).startswith("RGB"):
            return _hex(color.rgb)
        name = _THEME_COLOR_NAMES.get(str(color.theme_color).split()[0].split(".")[-1])
        return palette.get(name) if name else None
    except (AttributeError, TypeError, ValueError):
        return None


def _shape_fill_color(shape, palette: dict[str, str]) -> str | None:
    try:
        fill = shape.fill
        if fill.type != 1:  # MSO_FILL.SOLID
            return None
        color = fill.fore_color
        if str(color.type).startswith("RGB"):
            return _hex(color.rgb)
        name = _THEME_COLOR_NAMES.get(str(color.theme_color).split()[0].split(".")[-1])
        return palette.get(name) if name else None
    except (AttributeError, TypeError, ValueError, KeyError):
        return None


def probe(pptx_path: str | Path) -> MeasuredStyle:
    """Measure a template's style from whatever is actually in its slides."""
    try:
        prs = Presentation(str(pptx_path))
    except Exception as exc:  # noqa: BLE001 - python-pptx raises several types
        # Most often an Office lock file (`~$name.pptx`) picked up by a glob,
        # or a .ppt/.key renamed to .pptx.
        raise ValueError(
            f"{Path(pptx_path).name} is not a readable .pptx "
            f"({type(exc).__name__}). If it starts with '~$' it's a lock file "
            "PowerPoint created while the deck was open, not the deck itself."
        ) from exc
    palette = _theme_palette(prs)

    style = MeasuredStyle(
        slide_width_in=Emu(prs.slide_width).inches,
        slide_height_in=Emu(prs.slide_height).inches,
    )

    by_size: dict[float, Counter] = {}       # font size -> fonts used at that size
    color_by_size: dict[float, Counter] = {}  # font size -> colors used at that size
    chars_by_size: Counter = Counter()        # font size -> characters written at it
    slides_by_size: dict[float, set[int]] = {}  # font size -> which slides use it
    fills: Counter = Counter()
    lefts: list[float] = []
    rights: list[float] = []
    tops: list[float] = []
    bottoms: list[float] = []
    pictures: dict[tuple, list] = {}
    bottom_texts: Counter = Counter()
    n_slides = 0

    for slide in prs.slides:
        n_slides += 1
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            left, top = Emu(shape.left).inches, Emu(shape.top).inches
            right = left + Emu(shape.width).inches if shape.width else left
            bottom = top + Emu(shape.height).inches if shape.height else top

            if shape.shape_type == 13:  # PICTURE
                key = (round(left, 1), round(top, 1), round(right - left, 1))
                pictures.setdefault(key, []).append(shape)
                continue

            # Designed templates hang decorative shapes off the edges of the
            # slide. Measuring margins from those puts our content off-canvas,
            # so only shapes essentially inside the slide count.
            inside = (
                right > 0
                and bottom > 0
                and left > -0.05 * style.slide_width_in
                and right < 1.05 * style.slide_width_in
                and top > -0.05 * style.slide_height_in
                and bottom < 1.05 * style.slide_height_in
            )
            if inside:
                lefts.append(max(left, 0.0))
                rights.append(min(right, style.slide_width_in))
                tops.append(max(top, 0.0))
                bottoms.append(min(bottom, style.slide_height_in))

            fill = _shape_fill_color(shape, palette)
            if fill:
                fills[fill] += 1

            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    size = run.font.size.pt if run.font.size else None
                    if size is None:
                        continue
                    if run.font.name:
                        by_size.setdefault(size, Counter())[run.font.name] += 1
                    chars_by_size[size] += len(run.text or "")
                    slides_by_size.setdefault(size, set()).add(n_slides)
                    color = _run_color(run, palette)
                    if color:
                        color_by_size.setdefault(size, Counter())[color] += 1
                    # A footer is small text sitting low on the slide.
                    if size <= 12 and top > style.slide_height_in * 0.85 and text:
                        bottom_texts[text] += 1

    if not by_size and not lefts:
        style.notes.append("no measurable shapes found — defaults will be used throughout")
        return style

    sizes = sorted(by_size)
    if sizes:
        # Body text is whatever size carries the most actual text. Counting
        # runs instead would let a short footer repeated on every slide
        # outvote the real body copy. Restricted to sizes we also have a font
        # for: a run can specify a size and inherit its typeface, which would
        # otherwise pick a size absent from `by_size`.
        weighted = {size: n for size, n in chars_by_size.items() if size in by_size}
        body_size = max(weighted, key=lambda s: weighted[s]) if weighted else sizes[0]
        # A heading is a size bigger than the body that recurs across slides.
        # The plain largest size is unreliable — a one-off oversized figure
        # (a KPI number) would win it.
        bigger = [s for s in sizes if s > body_size]
        heading_size = (
            max(bigger, key=lambda s: (len(slides_by_size.get(s, ())), s)) if bigger else body_size
        )
        small = sizes[0]

        style.heading_font = by_size[heading_size].most_common(1)[0][0]
        style.body_font = by_size[body_size].most_common(1)[0][0]
        style.heading_pt, style.body_pt, style.small_pt = heading_size, body_size, small
        style.notes.append(
            f"fonts: heading={style.heading_font} (at {heading_size:g}pt), "
            f"body={style.body_font} (at {body_size:g}pt)"
        )

        if heading_size in color_by_size:
            style.colors["primary"] = color_by_size[heading_size].most_common(1)[0][0]
        if body_size in color_by_size:
            style.colors["text"] = color_by_size[body_size].most_common(1)[0][0]
        if small in color_by_size and small != body_size:
            style.colors["muted"] = color_by_size[small].most_common(1)[0][0]

    # The accent is the most-used genuinely colorful value, ignoring anything
    # already claimed as text/heading and anything grey.
    claimed = set(style.colors.values())
    candidates = Counter()
    for counter in list(color_by_size.values()) + [fills]:
        for value, count in counter.items():
            if value not in claimed and _colorfulness(value) > 40:
                candidates[value] += count
    if candidates:
        style.colors["accent"] = candidates.most_common(1)[0][0]

    # What was actually measured beats the theme, which is often left at the
    # stock Office scheme even in a deck with a deliberate palette.
    if style.colors.get("primary"):
        style.colors.setdefault("section_background", style.colors["primary"])
    for key, theme_key in (
        ("background", "lt1"), ("section_background", "dk2"),
        ("text", "dk1"), ("accent", "accent1"), ("primary", "dk2"),
    ):
        if palette.get(theme_key):
            style.colors.setdefault(key, palette[theme_key])

    if style.colors:
        style.notes.append("colors: " + ", ".join(f"{k}={v}" for k, v in sorted(style.colors.items())))

    if lefts:
        # Clamped to a plausible range: a measured margin of 0 (a full-bleed
        # band) or something enormous would ruin the layout either way.
        max_margin = 0.15 * style.slide_width_in
        style.margin_in = round(min(max(min(lefts), 0.3), max_margin), 2)
        style.content_right_in = round(min(max(rights), style.slide_width_in - style.margin_in), 2)
        style.title_top_in = round(min(max(min(tops), 0.3), 0.2 * style.slide_height_in), 2)
        style.content_bottom_in = round(
            min(max(bottoms), style.slide_height_in - 0.35), 2
        )
        style.notes.append(
            f"margins: left={style.margin_in}in, right edge={style.content_right_in}in, "
            f"first content at {style.title_top_in}in, last at {style.content_bottom_in}in"
        )

    # An image covering (nearly) the whole slide is the design's background,
    # not a logo. Reusing the file's own image is the only faithful way to
    # carry over a look that lives in artwork rather than in a fill colour —
    # a chalkboard, a textured paper, a photographic cover.
    slide_area = style.slide_width_in * style.slide_height_in
    full_bleed = [
        shapes[0]
        for shapes in pictures.values()
        if Emu(shapes[0].width).inches * Emu(shapes[0].height).inches > 0.85 * slide_area
    ]
    if full_bleed:
        try:
            style.background_image = max(
                full_bleed, key=lambda s: Emu(s.width).inches * Emu(s.height).inches
            ).image.blob
            style.notes.append(
                f"full-slide background image reused ({len(style.background_image) // 1024}KB)"
            )
        except (AttributeError, ValueError):
            pass

    # A logo is small, and repeats in the same spot. With a single slide
    # there's nothing to repeat, so fall back to a small image placed high up.
    def is_logo_sized(shape) -> bool:
        return (
            Emu(shape.width).inches <= 0.25 * style.slide_width_in
            and Emu(shape.height).inches <= 0.25 * style.slide_height_in
        )

    logo_shape = None
    repeated = [shapes for shapes in pictures.values() if len(shapes) > 1 and is_logo_sized(shapes[0])]
    if repeated:
        logo_shape = max(repeated, key=len)[0]
    elif n_slides == 1:
        singles = [s[0] for s in pictures.values() if is_logo_sized(s[0])]
        upper = [s for s in singles if Emu(s.top).inches < style.slide_height_in * 0.3]
        if upper:
            logo_shape = min(upper, key=lambda s: Emu(s.width).inches * Emu(s.height).inches)
    if logo_shape is not None:
        try:
            style.logo = logo_shape.image.blob
            style.logo_box_in = (
                round(Emu(logo_shape.left).inches, 2), round(Emu(logo_shape.top).inches, 2),
                round(Emu(logo_shape.width).inches, 2), round(Emu(logo_shape.height).inches, 2),
            )
            style.notes.append(f"logo found at {style.logo_box_in}")
        except (AttributeError, ValueError):
            style.logo = None

    # A footer is defined by recurring: it is the same line low on several
    # slides. Falling back to whatever single line sat at the bottom of one
    # slide picked up captions and, on a template that documents its own
    # palette, the swatch label « #FDDCC5 » — which was then printed across the
    # foot of every generated slide. The candidate also has to read as words,
    # and the one that recurs most often wins rather than the longest.
    repeated_footer = [
        text
        for text, count in bottom_texts.items()
        if count > 1 and sum(character.isalpha() for character in text) >= 3
    ]
    if repeated_footer:
        style.footer = max(repeated_footer, key=lambda t: (bottom_texts[t], len(t)))[:120]
        style.notes.append(f"footer: {style.footer!r}")

    return style
